"""Build initial DOCX templates programmatically using python-docx.

This script generates the initial template.docx for each template folder,
reproducing the layout from the sample PDFs as closely as possible.
Jinja2 tags are placed in single runs to avoid Word splitting them.

After generation, the club can edit these DOCX files in Word directly.
"""

import os
import sys
from pathlib import Path
from decimal import Decimal

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# Project root
ROOT = Path(__file__).parent.parent
ASSETS = ROOT / "assets"
TEMPLATES = ROOT / "templates"


def set_cell_border(cell, **kwargs):
    """Set cell border properties."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}></w:tcBorders>')
    for edge, val in kwargs.items():
        element = parse_xml(
            f'<w:{edge} {nsdecls("w")} w:val="{val.get("val", "single")}" '
            f'w:sz="{val.get("sz", "4")}" w:space="0" '
            f'w:color="{val.get("color", "000000")}"/>'
        )
        tcBorders.append(element)
    tcPr.append(tcBorders)


def set_cell_shading(cell, color: str):
    """Set cell background color."""
    shading = parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{color}" w:val="clear"/>'
    )
    cell._tc.get_or_add_tcPr().append(shading)


def add_paragraph_with_tag(doc_or_cell, text: str, style=None, bold=False,
                            font_size=None, font_name=None, alignment=None,
                            color=None, space_after=None, space_before=None):
    """Add a paragraph with text that may contain Jinja tags in single runs."""
    if hasattr(doc_or_cell, 'add_paragraph'):
        p = doc_or_cell.add_paragraph()
    else:
        p = doc_or_cell.paragraphs[0] if doc_or_cell.paragraphs else doc_or_cell.add_paragraph()
    
    if style:
        p.style = style
    if alignment is not None:
        p.alignment = alignment
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)

    run = p.add_run(text)
    if bold:
        run.bold = True
    if font_size:
        run.font.size = Pt(font_size)
    if font_name:
        run.font.name = font_name
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p


def add_header_logos(section, left_img=None, right_img=None, title_line1="", title_line2=""):
    """Add dual-logo header with centered text."""
    header = section.header
    header.is_linked_to_previous = False
    
    # Create a table for header layout: logo | text | logo
    htable = header.add_table(rows=1, cols=3, width=Inches(6.5))
    htable.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Left logo
    left_cell = htable.cell(0, 0)
    left_cell.width = Inches(1.0)
    if left_img and left_img.exists():
        lp = left_cell.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = lp.add_run()
        run.add_picture(str(left_img), width=Inches(0.8))
    
    # Center text
    center_cell = htable.cell(0, 1)
    center_cell.width = Inches(4.5)
    cp = center_cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if title_line1:
        run1 = cp.add_run(title_line1)
        run1.bold = True
        run1.font.size = Pt(14)
        run1.font.name = "Times New Roman"
        run1.font.color.rgb = RGBColor(0x1E, 0x5B, 0x94)
    if title_line2:
        run_br = cp.add_run()
        run_br.add_break()
        run2 = cp.add_run(title_line2)
        run2.font.size = Pt(9)
        run2.font.name = "Arial"
        run2.font.color.rgb = RGBColor(0x1E, 0x5B, 0x94)
    
    # Right logo
    right_cell = htable.cell(0, 2)
    right_cell.width = Inches(1.0)
    if right_img and right_img.exists():
        rp = right_cell.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = rp.add_run()
        run.add_picture(str(right_img), width=Inches(0.8))

    # Remove table borders
    for row in htable.rows:
        for cell in row.cells:
            set_cell_border(cell,
                top={"val": "none", "sz": "0"},
                bottom={"val": "none", "sz": "0"},
                start={"val": "none", "sz": "0"},
                end={"val": "none", "sz": "0"})


def add_svnit_header(section, header_img_name="clean_svnit_approval_header.png", crest_img=None):
    """Add SVNIT-style header using clean extracted logo+title band image from sample PDF."""
    header = section.header
    header.is_linked_to_previous = False
    
    clean_img = ASSETS / header_img_name
    if not clean_img.exists():
        clean_img = ASSETS / "clean_svnit_approval_header.png"
    if not clean_img.exists():
        clean_img = ASSETS / "svnit_bilingual_header.png"

    if clean_img.exists():
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hp.add_run()
        run.add_picture(str(clean_img), width=Inches(6.8))
    else:
        htable = header.add_table(rows=1, cols=3, width=Inches(6.5))
        htable.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Left: crest
        left_cell = htable.cell(0, 0)
        left_cell.width = Inches(1.2)
        if crest_img and crest_img.exists():
            lp = left_cell.paragraphs[0]
            lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = lp.add_run()
            run.add_picture(str(crest_img), width=Inches(0.9))
        
        # Center text
        center_cell = htable.cell(0, 1)
        center_cell.width = Inches(4.6)
        cp = center_cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        lines = [
            ("SARDAR VALLABHBHAI NATIONAL INSTITUTE OF TECHNOLOGY, SURAT", 11, True, "000000"),
            ("Ichchhanath, Surat - 395007 (Gujarat)", 9, False, "000000"),
        ]
        for i, (text, size, bold, color) in enumerate(lines):
            if i > 0:
                run = cp.add_run()
                run.add_break()
            run = cp.add_run(text)
            run.bold = bold
            run.font.size = Pt(size)
            run.font.name = "Times New Roman"
            run.font.color.rgb = RGBColor.from_string(color)

        # Right: ACM text/logo
        right_cell = htable.cell(0, 2)
        right_cell.width = Inches(0.7)
        rp = right_cell.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = rp.add_run("ACM")
        run.bold = True
        run.font.size = Pt(14)
        run.font.name = "Times New Roman"
        
        # Remove borders
        for row in htable.rows:
            for cell in row.cells:
                set_cell_border(cell,
                    top={"val": "none", "sz": "0"},
                    bottom={"val": "none", "sz": "0"},
                    start={"val": "none", "sz": "0"},
                    end={"val": "none", "sz": "0"})


# ====================================================================
# TEMPLATE 1: PERMISSION LETTER
# ====================================================================
def build_permission_letter():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)

    # Header: Single-line letterhead with both logos (cropped pure logo+title band, no REF NO or date)
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    perm_header_img = ASSETS / "clean_permission_header.png"
    if not perm_header_img.exists():
        perm_header_img = ASSETS / "acm_permission_header.png"
    if perm_header_img.exists():
        run_h = hp.add_run()
        run_h.add_picture(str(perm_header_img), width=Inches(6.8))
    else:
        svnit_crest = ASSETS / "permission_page1_img1.png"
        acm_logo = ASSETS / "permission_page1_img2.png"
        add_header_logos(
            section,
            left_img=svnit_crest,
            right_img=acm_logo,
            title_line1="ASSOCIATION for COMPUTING MACHINERY",
            title_line2="SARDAR VALLABHBHAI NATIONAL INSTITUTE OF TECHNOLOGY, SURAT",
        )

    # REF NO and DATE
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run("REF NO.: {{ ref_no }}")
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = "Times New Roman"
    run2 = p.add_run("\t\t\t\t\t\tDATE: {{ document_date }}")
    run2.bold = True
    run2.font.size = Pt(11)
    run2.font.name = "Times New Roman"

    # SUBMITTED TO
    add_paragraph_with_tag(doc, "SUBMITTED TO: {{ submitted_to }}", bold=True,
                           font_size=11, font_name="Times New Roman", space_after=4)

    # SUBJECT
    add_paragraph_with_tag(
        doc,
        "SUBJECT: Requesting Permission to Conduct {{ event_name }} in {{ venue }}.",
        bold=True, font_size=11, font_name="Times New Roman", space_after=4
    )

    # SCHEDULE
    add_paragraph_with_tag(doc, "SCHEDULE OF EVENT: {{ date_range_str }}",
                           bold=True, font_size=11, font_name="Times New Roman",
                           space_after=6)

    # ABOUT THE EVENT table
    table = doc.add_table(rows=6, cols=2, style='Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    header_cell = table.cell(0, 0)
    header_cell.merge(table.cell(0, 1))
    hp = header_cell.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run("ABOUT THE EVENT")
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = "Times New Roman"
    set_cell_shading(header_cell, "D9E2F3")

    rows_data = [
        ("Brief Description", "{{ description }}"),
        ("Objective", "{{ objectives }}"),
        ("Number of Participants", "{{ expected_participants }}"),
        ("Mode", "{{ mode }}"),
        ("Event Organizers", "{{ organizers }}"),
    ]
    for i, (label, value) in enumerate(rows_data, 1):
        cell0 = table.cell(i, 0)
        cell0.width = Inches(1.8)
        p0 = cell0.paragraphs[0]
        run = p0.add_run(label)
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"
        
        cell1 = table.cell(i, 1)
        cell1.width = Inches(4.9)
        p1 = cell1.paragraphs[0]
        run = p1.add_run(value)
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"

    doc.add_paragraph()  # spacing

    # BORDERED 2x2 Signature Grid matching sample PDF exactly
    sig_table = doc.add_table(rows=2, cols=2, style='Table Grid')
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Cell (0,0): Chairperson
    c00 = sig_table.cell(0, 0)
    p = c00.paragraphs[0]
    p.add_run("Name: {{ chairperson.name }}\n").font.name = "Times New Roman"
    p.add_run("Designation: Chairperson , ACM\n").font.name = "Times New Roman"
    p.add_run("Admission No.: {{ chairperson.admission_no }}\n").font.name = "Times New Roman"
    p.add_run("Branch: {{ chairperson.branch }}\n").font.name = "Times New Roman"
    p.add_run("Mobile No.: {% if chairperson.mobile %}{{ chairperson.mobile }}{% endif %}").font.name = "Times New Roman"  # Requirement 5: "Mobile" in Permission

    # Cell (0,1): Venue Tech Assistant
    c01 = sig_table.cell(0, 1)
    p = c01.paragraphs[0]
    p.add_run("Name: {% if venue_tech_assistant %}{{ venue_tech_assistant.name }}{% else %}{{ tech_assistant_name | default('') }}{% endif %}\n").font.name = "Times New Roman"
    p.add_run("Designation: {% if venue_tech_assistant and venue_tech_assistant.roles and venue_tech_assistant.roles.permission_letter %}{{ venue_tech_assistant.roles.permission_letter }}{% else %}Tech Assistant , {{ venue }}{% endif %}\n").font.name = "Times New Roman"
    p.add_run("Branch: Computer Science & Engineering\n").font.name = "Times New Roman"
    p.add_run("Mobile No.: {% if venue_tech_assistant and venue_tech_assistant.mobile %}{{ venue_tech_assistant.mobile }}{% endif %}").font.name = "Times New Roman"

    # Cell (1,0): Club Faculty Chairman
    c10 = sig_table.cell(1, 0)
    p = c10.paragraphs[0]
    p.add_run("Name: {{ faculty_chairman.name }}\n").font.name = "Times New Roman"
    p.add_run("Designation: Chairman , ACM\n").font.name = "Times New Roman"
    p.add_run("Department: Computer Science & Engineering\n").font.name = "Times New Roman"
    p.add_run("Mobile No.: {% if faculty_chairman.mobile %}{{ faculty_chairman.mobile }}{% endif %}").font.name = "Times New Roman"

    # Cell (1,1): Venue Faculty In-Charge
    c11 = sig_table.cell(1, 1)
    p = c11.paragraphs[0]
    p.add_run("Name: {% if venue_faculty_in_charge %}{{ venue_faculty_in_charge.name }}{% else %}{{ faculty_in_charge_name | default('') }}{% endif %}\n").font.name = "Times New Roman"
    p.add_run("Designation: {% if venue_faculty_in_charge and venue_faculty_in_charge.roles and venue_faculty_in_charge.roles.permission_letter %}{{ venue_faculty_in_charge.roles.permission_letter }}{% else %}Faculty In-charge , {{ venue }}{% endif %}\n").font.name = "Times New Roman"
    p.add_run("Department: Computer Science & Engineering\n").font.name = "Times New Roman"
    p.add_run("Mobile No.: {% if venue_faculty_in_charge and venue_faculty_in_charge.mobile %}{{ venue_faculty_in_charge.mobile }}{% endif %}").font.name = "Times New Roman"

    for r in sig_table.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)

    output = TEMPLATES / "permission_letter" / "template.docx"
    doc.save(str(output))
    print(f"  [OK] Permission Letter: {output}")


# ====================================================================
# TEMPLATE 2: PRIZE MONEY RECEIPT
# ====================================================================
def build_prize_money_receipt():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)

    # SVNIT clean logo+title header
    crest = ASSETS / "prize_page1_img1.png"
    add_svnit_header(section, header_img_name="clean_svnit_prize_header.png", crest_img=crest)

    # Title: Prize Money Receipt
    add_paragraph_with_tag(doc, "Prize Money Receipt", bold=True,
                           font_size=13, font_name="Times New Roman",
                           alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    # Sr. No. and Date line
    p = doc.add_paragraph()
    run = p.add_run("Sr. No. {{ sr_no }}")
    run.font.size = Pt(10.5)
    run.font.name = "Times New Roman"
    run2 = p.add_run("\t\t\t\t\t\tDate: {{ document_date }}")
    run2.font.size = Pt(10.5)
    run2.font.name = "Times New Roman"

    # Verbatim acknowledgment text matching sample PDF exactly
    add_paragraph_with_tag(
        doc,
        'It is hereby acknowledged that an amount of ₹ {{ winner_amount }} will be received by me for '
        'securing {{ winner_position }} Position in the competition/event named '
        '{{ event_name }} organized by {{ club_name }} held on {{ date_range_str }} at {{ venue }}.',
        font_size=10.5, font_name="Times New Roman", space_after=6
    )

    # (Account Detail) heading
    add_paragraph_with_tag(doc, "(Account Detail)", bold=True,
                           font_size=10.5, font_name="Times New Roman",
                           alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    # Account details table matching sample
    acct_table = doc.add_table(rows=7, cols=2, style='Table Grid')
    acct_data = [
        ("Affiliating Institute:", "{{ affiliating_institute_text }}"),
        ("Name as per Bank Account:", "{{ winner_bank_name_as_per_bank }}"),
        ("Bank Name:", "{{ winner_bank_name }}"),
        ("Branch:", "{{ winner_bank_branch }}"),
        ("Account No.:", "{{ winner_account_no }}"),
        ("IFSC Code:", "{{ winner_ifsc }}"),
        ("Pan Card No.:", "{{ winner_pan }}"),
    ]
    for i, (label, value) in enumerate(acct_data):
        cell0 = acct_table.cell(i, 0)
        p0 = cell0.paragraphs[0]
        run = p0.add_run(label)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.name = "Times New Roman"
        cell0.width = Inches(2.5)

        cell1 = acct_table.cell(i, 1)
        p1 = cell1.paragraphs[0]
        run = p1.add_run(value)
        run.font.size = Pt(9.5)
        run.font.name = "Times New Roman"

    doc.add_paragraph()

    # Student Details bullet block matching sample (Mobile No. per requirement 6)
    p_stu = doc.add_paragraph()
    p_stu.paragraph_format.space_after = Pt(4)
    run = p_stu.add_run("Student Details\n")
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.name = "Times New Roman"
    
    runs_stu = [
        "● Name: {{ winner_name }}\n",
        "● Roll No.: {{ winner_roll_no }}\n",
        "● Course & Semester: {{ winner_course_semester }}\n",
        "● Mobile No.: {{ winner_mobile }}\n",  # Rule 6: "Mobile No." in Prize Receipt
        "Signature:\n\n",
        "It is certified that the said amount is recommended for payment.\n\n",
        "Chairman Signature\n",
        "{{ faculty_chairman.name }}\n",
        "Chairman, ACM Student Chapter SVNIT",
    ]
    for line in runs_stu:
        r = p_stu.add_run(line)
        r.font.size = Pt(9.5)
        r.font.name = "Times New Roman"
        if "Chairman Signature" in line or "Name:" in line or "Roll No.:" in line:
            r.bold = True

    output = TEMPLATES / "prize_money_receipt" / "template.docx"
    doc.save(str(output))
    print(f"  [OK] Prize Money Receipt: {output}")


# ====================================================================
# TEMPLATE 3: APPROVAL NOTE
# ====================================================================
def build_approval_note():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)

    # Clean SVNIT header image (no ref no / date)
    crest = ASSETS / "approval_page1_img1.png"
    add_svnit_header(section, header_img_name="clean_svnit_approval_header.png", crest_img=crest)

    # Header line: No : and Date : exactly matching sample
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("No : {{ ref_no }}")
    run.font.size = Pt(9.5)
    run.font.name = "Times New Roman"
    run.bold = True
    run2 = p.add_run("\t\t\t\t\t\t\tDate : {{ document_date }}")
    run2.font.size = Pt(9.5)
    run2.font.name = "Times New Roman"
    run2.bold = True

    # Submitted to Dean & Subject exactly matching sample
    p_to = doc.add_paragraph()
    p_to.paragraph_format.space_after = Pt(2)
    r = p_to.add_run("Submitted to Dean (Students Welfare):\n")
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.name = "Times New Roman"
    r2 = p_to.add_run("Subject : Requesting Permission to Conduct {{ event_name }}")
    r2.bold = True
    r2.font.size = Pt(9.5)
    r2.font.name = "Times New Roman"

    # Meta Table
    subj_table = doc.add_table(rows=4, cols=2, style='Table Grid')
    subj_data = [
        ("Chapter/Club/Society/Section Name", ": {{ club_name }}"),
        ("Chapter/Club/Society/Section Code No", ": {{ chapter_code }}"),
        ("Schedule of Event (Date, Day)", ": {{ date_range_str }}"),
        ("Event Venue", ": {{ venue }}"),
    ]
    for i, (label, value) in enumerate(subj_data):
        cell0 = subj_table.cell(i, 0)
        p0 = cell0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        run = p0.add_run(label)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"
        cell0.width = Inches(2.6)

        cell1 = subj_table.cell(i, 1)
        p1 = cell1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        run = p1.add_run(value)
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"

    # About the Event table
    add_paragraph_with_tag(doc, "About the Event", bold=True,
                           font_size=9.5, font_name="Times New Roman", space_before=1, space_after=1)

    about_table = doc.add_table(rows=7, cols=2, style='Table Grid')
    about_rows = [
        ("Title of the Event", ": {{ event_name }}"),
        ("Category of the Event", ": {{ category }}"),
        ("Brief Description", ": {{ description }}"),
        ("Objectives", ": {{ objectives }}"),
        ("Target Audience", ": {{ target_audience }}"),
        ("Expected Number of Participants", ": {{ expected_participants }}"),
        ("Name of Judges", ": {% if judges %}{% for judge in judges %}{{ judge }}{% if not loop.last %}, {% endif %}{% endfor %}{% else %}Nil{% endif %}"),
    ]
    for i, (label, value) in enumerate(about_rows):
        cell0 = about_table.cell(i, 0)
        p0 = cell0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        run = p0.add_run(label)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"
        cell0.width = Inches(2.6)

        cell1 = about_table.cell(i, 1)
        p1 = cell1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        run = p1.add_run(value)
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"

    # Financial Proposal
    add_paragraph_with_tag(doc, "Financial Proposal (Event Expenses)", bold=True,
                           font_size=9.5, font_name="Times New Roman", space_before=1, space_after=1)

    fin_table = doc.add_table(rows=1, cols=2, style='Table Grid')
    cell_h0 = fin_table.cell(0, 0)
    p = cell_h0.paragraphs[0]
    p.add_run("Items").bold = True
    cell_h0.width = Inches(4.5)

    cell_h1 = fin_table.cell(0, 1)
    p = cell_h1.paragraphs[0]
    p.add_run("Approximate Amount (Rs)").bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    cell_h1.width = Inches(2.5)

    for cell in [cell_h0, cell_h1]:
        set_cell_shading(cell, "D9E2F3")
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        cell.paragraphs[0].runs[0].font.name = "Times New Roman"

    # Budget items loop: start row
    r_start = fin_table.add_row()
    r_start.cells[0].paragraphs[0].text = "{%tr for item in budget_items %}"

    # Data row
    row = fin_table.add_row()
    cell0 = row.cells[0]
    p0 = cell0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    run = p0.add_run("{{ item.item }}")
    run.font.size = Pt(9)
    run.font.name = "Times New Roman"

    cell1 = row.cells[1]
    p1 = cell1.paragraphs[0]
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p1.add_run("{{ item.amount }}")
    run.font.size = Pt(9)
    run.font.name = "Times New Roman"

    # End row
    r_end = fin_table.add_row()
    r_end.cells[0].paragraphs[0].text = "{%tr endfor %}"

    # Total row
    total_row = fin_table.add_row()
    tc0 = total_row.cells[0]
    p = tc0.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("Total")
    run.bold = True
    run.font.size = Pt(9)
    run.font.name = "Times New Roman"

    tc1 = total_row.cells[1]
    p = tc1.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("{{ budget_total_inr }}")
    run.bold = True
    run.font.size = Pt(9)
    run.font.name = "Times New Roman"

    # Office Use block (matching sample PDF verbatim)
    add_paragraph_with_tag(doc, "Office Use", bold=True,
                           font_size=9.5, font_name="Times New Roman", space_before=1, space_after=1)

    off_table = doc.add_table(rows=5, cols=2, style='Table Grid')
    off_data = [
        ("Total amount Sanctioned in Current Academic Year", ": "),
        ("Amount Utilized", ": "),
        ("Current Proposed amount", ": {{ budget_total_inr }}"),
        ("Balance amount", ": "),
        ("Signature (Office Clerk)", ": "),
    ]
    for i, (label, val) in enumerate(off_data):
        c0 = off_table.cell(i, 0)
        c0.width = Inches(3.5)
        p = c0.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        r.font.size = Pt(9)
        r.font.name = "Times New Roman"
        r.bold = True

        c1 = off_table.cell(i, 1)
        c1.width = Inches(3.5)
        p = c1.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(val)
        r.font.size = Pt(9)
        r.font.name = "Times New Roman"

    # Recommended/Not Recommended & Approved lines matching sample exactly
    p_rec = doc.add_paragraph()
    p_rec.paragraph_format.space_before = Pt(1)
    p_rec.paragraph_format.space_after = Pt(0)
    r1 = p_rec.add_run("Recommended/Not Recommended\t\tRecommended/Not Recommended\t\tApproved/Not Approved")
    r1.font.size = Pt(9)
    r1.font.name = "Times New Roman"
    r1.bold = True

    p_names = doc.add_paragraph()
    p_names.paragraph_format.space_before = Pt(2)
    p_names.paragraph_format.space_after = Pt(1)
    r2 = p_names.add_run("Dr. Rakesh Maurya\t\t\tDr. Manish K Rathod\t\t\tDr. Sanjay R. Patel\n")
    r2.font.size = Pt(9)
    r2.font.name = "Times New Roman"
    r2.bold = True
    r3 = p_names.add_run("Chairman, F& SW\t\t\tAssociate Dean(SW)\t\t\tDean (Students Welfare)")
    r3.font.size = Pt(9)
    r3.font.name = "Times New Roman"

    # 4 Signature blocks matching sample exactly (Moblie No. per requirement 6)
    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sigs = [
        ("Signature of Student Head\nName: {{ student_head.name }}\nRoll No. {{ student_head.admission_no }}\nBranch: {{ student_head.branch }}\nMobile No. {{ student_head.mobile }}",
         "Signature of Student Secretary\nName: {{ student_secretary.name }}\nRoll No. {{ student_secretary.admission_no }}\nBranch: {{ student_secretary.branch }}\nMobile No. {{ student_secretary.mobile }}"),
        ("Signature of Faculty Co-Chairman\nName:           {{ faculty_co_chairman.name }}\nDesignation: Co-Chairman\nDept:            CSE\nMoblie No.   {{ faculty_co_chairman.mobile }}",
         "Signature of Faculty Chairman\nName:           {{ faculty_chairman.name }}\nDesignation: Chairman\nDept:             CSE\nMoblie No.   {{ faculty_chairman.mobile }}"),
    ]
    for ri, (left, right) in enumerate(sigs):
        lcell = sig_table.cell(ri, 0)
        lp = lcell.paragraphs[0]
        lp.paragraph_format.space_before = Pt(0)
        lp.paragraph_format.space_after = Pt(0)
        run = lp.add_run(left)
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"

        rcell = sig_table.cell(ri, 1)
        rp = rcell.paragraphs[0]
        rp.paragraph_format.space_before = Pt(0)
        rp.paragraph_format.space_after = Pt(0)
        run = rp.add_run(right)
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"

        for cell in [lcell, rcell]:
            set_cell_border(cell,
                top={"val": "none", "sz": "0"},
                bottom={"val": "none", "sz": "0"},
                start={"val": "none", "sz": "0"},
                end={"val": "none", "sz": "0"})

    # Ensure rows don't split across pages
    for tbl in [subj_table, about_table, fin_table, off_table, sig_table]:
        for r in tbl.rows:
            trPr = r._tr.get_or_add_trPr()
            trPr.append(parse_xml(r'<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

    output = TEMPLATES / "approval_note" / "template.docx"
    doc.save(str(output))
    print(f"  [OK] Approval Note: {output}")


# ====================================================================
# TEMPLATE 4: BILL SUMMARY
# ====================================================================
def build_bill_summary():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)

    # ACM text header (this sample uses text, not logos)
    add_paragraph_with_tag(doc, "ACM Student Chapter", bold=True,
                           font_size=14, font_name="Times New Roman",
                           alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_paragraph_with_tag(
        doc,
        "Sardar Vallabhbhai National Institute of Technology, Surat",
        font_size=10, font_name="Times New Roman",
        alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=8
    )

    # Date
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Date: {{ document_date }}")
    run.font.size = Pt(10)
    run.font.name = "Times New Roman"

    # Title
    add_paragraph_with_tag(
        doc,
        "Reimbursement for Cost incurred in {{ event_name }}",
        bold=True, font_size=12, font_name="Times New Roman",
        alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=8
    )

    # Table 1: Main bill summary
    t1 = doc.add_table(rows=1, cols=7, style='Table Grid')
    headers = ["Sr. No.", "Name", "Account Details", "To Be Paid to",
               "Prize Money (INR)", "Other", "Grand Total"]
    for i, h in enumerate(headers):
        cell = t1.cell(0, i)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(8)
        run.font.name = "Times New Roman"
        set_cell_shading(cell, "D9E2F3")

    # Winner rows: start row
    r_start1 = t1.add_row()
    r_start1.cells[0].paragraphs[0].text = "{%tr for winner in winners %}"

    row = t1.add_row()
    cells = row.cells
    for i, tag in enumerate([
        "{{ loop.index }}",
        "{{ winner.name }}",
        "{{ winner.bank.account_no }}",
        "{{ winner.bank.name_as_per_bank }}",
        "{{ winner.amount }}",
        "",
        "{{ winner.amount }}"
    ]):
        p = cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(tag)
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"

    # End row
    r_end1 = t1.add_row()
    r_end1.cells[0].paragraphs[0].text = "{%tr endfor %}"

    # Grand total row
    total_row = t1.add_row()
    merged = total_row.cells[0]
    for i in range(1, 4):
        merged.merge(total_row.cells[i])
    p = merged.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Grand Total")
    run.bold = True
    run.font.size = Pt(8.5)
    run.font.name = "Times New Roman"

    # Prize Money Total
    p4 = total_row.cells[4].paragraphs[0]
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run4 = p4.add_run("{{ prize_total_inr }}")
    run4.bold = True
    run4.font.size = Pt(8.5)
    run4.font.name = "Times New Roman"

    # Other Total
    p5 = total_row.cells[5].paragraphs[0]
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run5 = p5.add_run("{{ other_total_inr }}")
    run5.bold = True
    run5.font.size = Pt(8.5)
    run5.font.name = "Times New Roman"

    # Overall Grand Total
    p6 = total_row.cells[6].paragraphs[0]
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run6 = p6.add_run("{{ grand_total_inr }}")
    run6.bold = True
    run6.font.size = Pt(8.5)
    run6.font.name = "Times New Roman"

    doc.add_paragraph()

    # Signatures for Table 1 exactly matching sample
    sig1 = doc.add_table(rows=1, cols=3)
    for cell in sig1.rows[0].cells:
        set_cell_border(cell,
            top={"val": "none", "sz": "0"},
            bottom={"val": "none", "sz": "0"},
            start={"val": "none", "sz": "0"},
            end={"val": "none", "sz": "0"})
    
    sigs = [
        "Student Head\n{{ student_head.name }}\nACM NIT-Surat",
        "Student Secretary\n{{ student_secretary.name }}\nACM NIT-Surat",
        "Faculty Chairman\n{{ faculty_chairman.name }}\nACM NIT-Surat",
    ]
    for i, sig in enumerate(sigs):
        p = sig1.cell(0, i).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"\n\n\n{sig}")
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"

    doc.add_paragraph()

    # Table 2: Other Bank Account Details
    add_paragraph_with_tag(doc, "Other Bank Account Details", bold=True,
                           font_size=11, font_name="Times New Roman",
                           alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    t2 = doc.add_table(rows=1, cols=5, style='Table Grid')
    headers2 = ["Sr. No.", "Name", "Account No.", "IFSC Code", "Amount"]
    for i, h in enumerate(headers2):
        cell = t2.cell(0, i)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"
        set_cell_shading(cell, "D9E2F3")

    # Other reimbursement rows: start row
    r_start2 = t2.add_row()
    r_start2.cells[0].paragraphs[0].text = "{%tr for reimb in other_reimbursements %}"

    row = t2.add_row()
    for i, tag in enumerate([
        "{{ loop.index }}",
        "{{ reimb.name }}",
        "{{ reimb.account_no }}",
        "{{ reimb.ifsc }}",
        "{{ reimb.amount }}"
    ]):
        p = row.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(tag)
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"

    # End row
    r_end2 = t2.add_row()
    r_end2.cells[0].paragraphs[0].text = "{%tr endfor %}"

    # Total row
    total_row = t2.add_row()
    merged = total_row.cells[0]
    for i in range(1, 4):
        merged.merge(total_row.cells[i])
    p = merged.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Total:")
    run.bold = True
    run.font.size = Pt(9)
    run.font.name = "Times New Roman"

    p_tot2 = total_row.cells[4].paragraphs[0]
    p_tot2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_tot2 = p_tot2.add_run("{{ other_total_inr }}")
    run_tot2.bold = True
    run_tot2.font.size = Pt(9)
    run_tot2.font.name = "Times New Roman"

    doc.add_paragraph()

    # Faculty Chairman signature under table 2
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("\n\n\nFaculty Chairman\n{{ faculty_chairman.name }}")
    run.font.size = Pt(9)
    run.font.name = "Times New Roman"

    output = TEMPLATES / "bill_summary" / "template.docx"
    doc.save(str(output))
    print(f"  [OK] Bill Summary: {output}")


# ====================================================================
# TEMPLATE 5: BILL CERTIFICATE (FORM BC-R)
# ====================================================================
def build_bill_certificate():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
def build_bill_certificate():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)

    # Clean SVNIT Bill Certificate header image
    crest = ASSETS / "billcert_page1_img1.png"
    add_svnit_header(section, header_img_name="clean_svnit_billcert_header.png", crest_img=crest)

    # Header Titles matching sample PDF exactly (No invented subtitle!)
    p_hdr = doc.add_paragraph()
    p_hdr.paragraph_format.space_before = Pt(2)
    p_hdr.paragraph_format.space_after = Pt(2)
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_hdr.add_run("DEAN STUDENT WELFARE\nFORM BC-R\nBILL CERTIFICATE\n(RECURRING EXPENDITURE)")
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = "Times New Roman"

    # Meta lines
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(2)
    p_meta.paragraph_format.space_after = Pt(2)
    r_m = p_meta.add_run(
        "Chapter Name:                                      {{ club_name }}                                       Code No .        {{ chapter_code }}\n"
        "Approval Note No.:  {{ approval_note_no }}                  Date:   {{ approval_note_date }}\n"
        "Original Approval Note/Comparative Statement attached herewith/sent with bill No. {{ invoice_no }}\n"
        "dtd.: {{ invoice_date }}                                   for Rs. {{ bill_total_with_tax }}                          of M/s. {{ supplier_name }}\n"
        "Purchase Order No. and date : {{ purchase_order }}\n"
        "Head of Account (Code No.) : {{ head_of_account }}\n"
        "Advance drawn in Rs. (if any) : {{ advance }}"
    )
    r_m.font.size = Pt(8.5)
    r_m.font.name = "Times New Roman"

    # Certification statement verbatim from sample (including "I, am")
    p_cert = doc.add_paragraph()
    p_cert.paragraph_format.space_before = Pt(2)
    p_cert.paragraph_format.space_after = Pt(3)
    r_c = p_cert.add_run(
        '“I, am {% if certifier_name %}{{ certifier_name }}{% else %}____________________{% endif %}, '
        'personally satisfied that the goods (described below) purchased are of the requisite quality and '
        'specification and have been purchased from a reliable supplier / contractor at a reasonable price.”'
    )
    r_c.font.size = Pt(8.5)
    r_c.font.name = "Times New Roman"

    # Items table matching sample
    items_table = doc.add_table(rows=1, cols=5, style='Table Grid')
    headers_it = ["Sr. No.", "Name of the Equipment/ item(s)", "Quantity", "Unit Cost (Rs.)", "Total Cost of Item(s) (Rs.)"]
    for i, h in enumerate(headers_it):
        cell = items_table.cell(0, i)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(8.0)
        run.font.name = "Times New Roman"
        set_cell_shading(cell, "D9E2F3")

    # Item rows: start row
    r_start = items_table.add_row()
    r_start.cells[0].paragraphs[0].text = "{%tr for item in bill_items %}"

    row = items_table.add_row()
    for i, tag in enumerate([
        "{{ loop.index }}",
        "{{ item.name }}",
        "{{ item.qty }}",
        "{{ item.unit_cost }}",
        "{{ item.line_total }}"
    ]):
        p = row.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(tag)
        run.font.size = Pt(8.0)
        run.font.name = "Times New Roman"

    # End row
    r_end = items_table.add_row()
    r_end.cells[0].paragraphs[0].text = "{%tr endfor %}"

    # Totals
    for label, val in [
        ("Total Amount (Rs.)", "{{ bill_total }}"),
        ("+ Taxes or other charges", "{{ tax_amount }}"),
        ("Total Estimated Amount Including Taxes (Rs.)", "{{ bill_total_with_tax }}")
    ]:
        row = items_table.add_row()
        merged = row.cells[0]
        for c in range(1, 4):
            merged.merge(row.cells[c])
        p = merged.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(8.0)
        r.font.name = "Times New Roman"

        p_v = row.cells[4].paragraphs[0]
        p_v.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_v = p_v.add_run(val)
        r_v.bold = True
        r_v.font.size = Pt(8.0)
        r_v.font.name = "Times New Roman"

    # SUPPLIER DETAILS header
    add_paragraph_with_tag(doc, "SUPPLIER DETAILS", bold=True,
                           font_size=8.5, font_name="Times New Roman", space_before=2, space_after=1)

    sup_table = doc.add_table(rows=4, cols=2, style='Table Grid')
    sup_data = [
        ("Name of the Supplier", "{{ supplier_name }}"),
        ("GST No.", "{{ supplier_gst }}"),
        ("Invoice No. and date", "{{ invoice_no }} and {{ invoice_date }}"),
        ("Amount (Rs.)", "{{ bill_total_with_tax }}"),
    ]
    for i, (label, value) in enumerate(sup_data):
        cell0 = sup_table.cell(i, 0)
        p = cell0.paragraphs[0]
        run = p.add_run(label)
        run.bold = True
        run.font.size = Pt(8.0)
        run.font.name = "Times New Roman"
        cell0.width = Inches(2.5)

        cell1 = sup_table.cell(i, 1)
        p = cell1.paragraphs[0]
        run = p.add_run(value)
        run.font.size = Pt(8.0)
        run.font.name = "Times New Roman"

    # PAYMENT DETAILS header
    add_paragraph_with_tag(doc, "PAYMENT DETAILS", bold=True,
                           font_size=8.5, font_name="Times New Roman", space_before=2, space_after=1)

    pay_table = doc.add_table(rows=6, cols=2, style='Table Grid')
    pay_data = [
        ("Name of the Party", "{{ payment_party }}"),
        ("Account Number", "{{ payment_account_no }}"),
        ("Account Holder Name", "{{ payment_holder }}"),
        ("Name of the Bank", "{{ payment_bank }}"),
        ("Name of the Branch", "{{ payment_branch }}"),
        ("IFSC Code", "{{ payment_ifsc }}"),
    ]
    for i, (label, value) in enumerate(pay_data):
        cell0 = pay_table.cell(i, 0)
        p = cell0.paragraphs[0]
        run = p.add_run(label)
        run.bold = True
        run.font.size = Pt(8.0)
        run.font.name = "Times New Roman"
        cell0.width = Inches(2.5)

        cell1 = pay_table.cell(i, 1)
        p = cell1.paragraphs[0]
        run = p.add_run(value)
        run.font.size = Pt(8.0)
        run.font.name = "Times New Roman"

    # --- EXACTLY PAGE 2 (Requirement 7) ---
    doc.add_page_break()

    # STOCK ENTRY DETAILS matching sample exactly
    p_stk_hdr = doc.add_paragraph()
    p_stk_hdr.paragraph_format.space_before = Pt(4)
    p_stk_hdr.paragraph_format.space_after = Pt(2)
    r = p_stk_hdr.add_run("STOCK ENTRY DETAILS\nName of the Laboratory / Office: {{ laboratory_office }}")
    r.bold = True
    r.font.size = Pt(9.0)
    r.font.name = "Times New Roman"

    stock_table = doc.add_table(rows=6, cols=4, style='Table Grid')
    stock_headers = ["Name/Type", "Recurring Register", "Consumable Register", "Miscellaneous Register"]
    for i, h in enumerate(stock_headers):
        c = stock_table.cell(0, i)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(8.0)
        r.font.name = "Times New Roman"
        set_cell_shading(c, "D9E2F3")

    stock_rows_data = [
        ("Serial Number", "", "", ""),
        ("Page Number", "", "", ""),
        ("Date", "", "", ""),
        ("Name and Signature of a Clerk/Staff", "", "", ""),
    ]
    for ri, rdata in enumerate(stock_rows_data, 1):
        for ci, val in enumerate(rdata):
            c = stock_table.cell(ri, ci)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.0)
            r.font.name = "Times New Roman"
            if ci == 0:
                r.bold = True

    # INDENTER DETAILS
    add_paragraph_with_tag(doc, "INDENTER DETAILS", bold=True,
                           font_size=9.0, font_name="Times New Roman", space_before=4, space_after=2)

    ind_table = doc.add_table(rows=4, cols=2, style='Table Grid')
    ind_data = [
        ("Name of the Faculty Chairman/Co-Chairman", "{{ faculty_chairman.name }}"),
        ("Designation", "Chairman, ACM Student Chapter SVNIT"),
        ("Department", "CSE"),
        ("Signature with Date", "\n\n"),
    ]
    for i, (label, value) in enumerate(ind_data):
        cell0 = ind_table.cell(i, 0)
        p = cell0.paragraphs[0]
        run = p.add_run(label)
        run.bold = True
        run.font.size = Pt(8.0)
        run.font.name = "Times New Roman"
        cell0.width = Inches(2.5)

        cell1 = ind_table.cell(i, 1)
        p = cell1.paragraphs[0]
        run = p.add_run(value)
        run.font.size = Pt(8.0)
        run.font.name = "Times New Roman"

    # RECOMMENDATIONS (HoD & Dean) matching sample PDF exactly
    add_paragraph_with_tag(doc, "Recommendation (Head of the Department)", bold=True,
                           font_size=9.0, font_name="Times New Roman", space_before=4, space_after=2)
    add_paragraph_with_tag(doc, "Recommended Amount: {{ bill_total_with_tax }}\nApproved and Recommended for Bill Payment\n\n\nSignature of HoD (CSE)",
                           font_size=8.5, font_name="Times New Roman", space_after=4)

    add_paragraph_with_tag(doc, "Dean (Students Welfare) Approval", bold=True,
                           font_size=9.0, font_name="Times New Roman", space_before=4, space_after=2)
    add_paragraph_with_tag(doc, "Approved / Not Approved\n\n\nDean (Students Welfare)",
                           font_size=8.5, font_name="Times New Roman", space_after=4)

    output = TEMPLATES / "bill_certificate" / "template.docx"
    doc.save(str(output))
    print(f"  [OK] Bill Certificate: {output}")


# ====================================================================
# TEMPLATE 6: MINUTES OF MEETING (PLACEHOLDER)
# ====================================================================
def build_mom():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)

    # Header: Single-line letterhead with both logos (clean pure logo+title band, no wrapping)
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mom_header_img = ASSETS / "clean_permission_header.png"
    if not mom_header_img.exists():
        mom_header_img = ASSETS / "acm_permission_header.png"
    if mom_header_img.exists():
        run_h = hp.add_run()
        run_h.add_picture(str(mom_header_img), width=Inches(6.8))
    else:
        svnit_crest = ASSETS / "permission_page1_img1.png"
        acm_logo = ASSETS / "permission_page1_img2.png"
        add_header_logos(
            section,
            left_img=svnit_crest,
            right_img=acm_logo,
            title_line1="ASSOCIATION for COMPUTING MACHINERY",
            title_line2="SARDAR VALLABHBHAI NATIONAL INSTITUTE OF TECHNOLOGY, SURAT",
        )


    # Title
    add_paragraph_with_tag(doc, "MINUTES OF MEETING", bold=True,
                           font_size=14, font_name="Times New Roman",
                           alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    # Ref and Date
    p = doc.add_paragraph()
    run = p.add_run("Ref No.: {{ ref_no }}")
    run.font.size = Pt(11)
    run.font.name = "Times New Roman"
    run2 = p.add_run("\t\t\t\t\tDate: {{ meeting_date }}")
    run2.font.size = Pt(11)
    run2.font.name = "Times New Roman"

    # Meeting details table
    details = doc.add_table(rows=5, cols=2, style='Table Grid')
    detail_data = [
        ("Meeting Title", "{{ meeting_title }}"),
        ("Date & Time", "{{ meeting_date }}, {{ meeting_start_time }} - {{ meeting_end_time }}"),
        ("Venue", "{{ venue }}"),
        ("Mode", "{{ mode }}"),
        ("Chaired By", "{{ chaired_by }}"),
    ]
    for i, (label, value) in enumerate(detail_data):
        cell0 = details.cell(i, 0)
        p = cell0.paragraphs[0]
        run = p.add_run(label)
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"
        cell0.width = Inches(2.0)

        cell1 = details.cell(i, 1)
        p = cell1.paragraphs[0]
        run = p.add_run(value)
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"

    doc.add_paragraph()

    # Attendees
    add_paragraph_with_tag(doc, "Attendees:", bold=True,
                           font_size=11, font_name="Times New Roman", space_after=2)
    add_paragraph_with_tag(
        doc,
        "{% for person in attendees %}{{ loop.index }}. {{ person }}\n{% endfor %}",
        font_size=10, font_name="Times New Roman", space_after=4
    )

    # Absentees
    add_paragraph_with_tag(doc, "Absentees:", bold=True,
                           font_size=11, font_name="Times New Roman", space_after=2)
    add_paragraph_with_tag(
        doc,
        "{% if absentees and absentees|length > 0 and absentees[0] != 'None' and absentees[0] != None %}{% for person in absentees %}{{ loop.index }}. {{ person }}\n{% endfor %}{% else %}Nil\n{% endif %}",
        font_size=10, font_name="Times New Roman", space_after=4
    )

    # Agenda & Discussion
    add_paragraph_with_tag(doc, "Agenda & Discussion:", bold=True,
                           font_size=11, font_name="Times New Roman", space_after=4)

    agenda_table = doc.add_table(rows=1, cols=4, style='Table Grid')
    for i, h in enumerate(["Sr.", "Agenda Item", "Discussion", "Decision"]):
        cell = agenda_table.cell(0, i)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"
        set_cell_shading(cell, "D9E2F3")

    # Agenda items loop: start row
    r_start_ag = agenda_table.add_row()
    r_start_ag.cells[0].paragraphs[0].text = "{%tr for item in agenda_items %}"

    row = agenda_table.add_row()
    for i, tag in enumerate([
        "{{ loop.index }}",
        "{{ item.topic }}",
        "{{ item.discussion }}",
        "{{ item.decision }}"
    ]):
        p = row.cells[i].paragraphs[0]
        run = p.add_run(tag)
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"

    # End row
    r_end_ag = agenda_table.add_row()
    r_end_ag.cells[0].paragraphs[0].text = "{%tr endfor %}"

    doc.add_paragraph()

    # Action Items
    add_paragraph_with_tag(doc, "Action Items:", bold=True,
                           font_size=11, font_name="Times New Roman", space_after=4)

    action_table = doc.add_table(rows=1, cols=4, style='Table Grid')
    for i, h in enumerate(["Sr.", "Task", "Owner", "Deadline"]):
        cell = action_table.cell(0, i)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"
        set_cell_shading(cell, "D9E2F3")

    # Action items loop: start row
    r_start_ac = action_table.add_row()
    r_start_ac.cells[0].paragraphs[0].text = "{%tr for item in action_items %}"

    row = action_table.add_row()
    for i, tag in enumerate([
        "{{ loop.index }}",
        "{{ item.task }}",
        "{{ item.owner }}",
        "{{ item.deadline }}"
    ]):
        p = row.cells[i].paragraphs[0]
        run = p.add_run(tag)
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"

    # End row
    r_end_ac = action_table.add_row()
    r_end_ac.cells[0].paragraphs[0].text = "{%tr endfor %}"

    doc.add_paragraph()

    # Next meeting
    add_paragraph_with_tag(doc, "Next Meeting: {{ next_meeting_date }}",
                           font_size=10, font_name="Times New Roman", space_after=8)

    # Signatures
    sig_table = doc.add_table(rows=1, cols=2)
    for cell in sig_table.rows[0].cells:
        set_cell_border(cell,
            top={"val": "none", "sz": "0"},
            bottom={"val": "none", "sz": "0"},
            start={"val": "none", "sz": "0"},
            end={"val": "none", "sz": "0"})

    p = sig_table.cell(0, 0).paragraphs[0]
    run = p.add_run("\n\n\nPrepared by:\n{{ prepared_by }}")
    run.font.size = Pt(9)
    run.font.name = "Times New Roman"

    p = sig_table.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("\n\n\nApproved by:\n{{ approved_by }}")
    run.font.size = Pt(9)
    run.font.name = "Times New Roman"

    output = TEMPLATES / "mom" / "template.docx"
    doc.save(str(output))
    print(f"  [OK] MoM (placeholder): {output}")


# ====================================================================
# MAIN
# ====================================================================
def main():
    print("Building DOCX templates from sample PDF analysis...")
    print(f"Assets dir: {ASSETS}")
    print(f"Templates dir: {TEMPLATES}")
    print()

    build_permission_letter()
    build_prize_money_receipt()
    build_approval_note()
    build_bill_summary()
    build_bill_certificate()
    build_mom()

    print("\nAll templates built! [OK]")
    print("Run 'python -m docgen validate-templates' to validate.")


if __name__ == "__main__":
    main()
