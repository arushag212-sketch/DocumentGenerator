"""Image blocks in editor templates: chosen images print, and an empty club-logo block prints
the rendering club's own logo."""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document

from docgen.config import ClubConfig
from docgen.letterhead_banner import LOGO_PLACEHOLDER_NAME
from docgen.renderer import Renderer
from docxtpl import DocxTemplate
from server.app.core.template_compiler import TemplateCompiler

LOGO = Path("assets/clubs/raag.png")


def _compile(blocks: list[dict]) -> bytes:
    docx_bytes, _schema, report = TemplateCompiler.compile(
        editor_json={"type": "doc", "content": blocks},
        fields=[],
        template_meta={"id": "img_test", "name": "Image Test"},
        club_slug="test",
        letterhead_cfg={"banner_placeholder": True},
    )
    assert report["is_valid"], report["errors"]
    return docx_bytes


def _body_media(path: Path) -> list[bytes]:
    with zipfile.ZipFile(path) as z:
        body_rels = z.read("word/_rels/document.xml.rels").decode()
        return [z.read("word/" + n.split('Target="')[1].split('"')[0])
                for n in body_rels.split("<Relationship ")[1:] if "media/" in n]


def test_empty_club_logo_block_prints_the_clubs_logo(tmp_path):
    src = tmp_path / "t.docx"
    src.write_bytes(_compile([{"type": "imageBlock", "attrs": {"imageType": "club-logo", "src": ""}},
                              {"type": "paragraph", "content": [{"type": "text", "text": "Body"}]}]))
    assert LOGO_PLACEHOLDER_NAME in "".join(p._element.xml for p in Document(str(src)).paragraphs)

    doc = DocxTemplate(str(src))
    Renderer(config=ClubConfig(club_name="Raag, SVNIT", club_logo_path=str(LOGO)))._apply_club_logo(doc)
    doc.render({}, autoescape=True)
    out = tmp_path / "out.docx"
    doc.save(str(out))
    assert LOGO.read_bytes() in _body_media(out)


def test_empty_club_logo_block_is_dropped_when_club_has_no_logo(tmp_path):
    src = tmp_path / "t.docx"
    src.write_bytes(_compile([{"type": "imageBlock", "attrs": {"imageType": "club-logo", "src": ""}}]))
    doc = DocxTemplate(str(src))
    Renderer(config=ClubConfig(club_name="No Logo Club"))._apply_club_logo(doc)
    doc.render({}, autoescape=True)
    out = tmp_path / "out.docx"
    doc.save(str(out))
    with zipfile.ZipFile(out) as z:
        body = z.read("word/document.xml").decode()
    assert "<w:drawing>" not in body and LOGO_PLACEHOLDER_NAME not in body, "no grey placeholder box on the page"


def test_uploaded_image_block_prints(tmp_path):
    import base64

    data_uri = "data:image/png;base64," + base64.b64encode(LOGO.read_bytes()).decode()
    src = tmp_path / "t.docx"
    src.write_bytes(_compile([{"type": "imageBlock", "attrs": {"imageType": "image", "src": data_uri, "width": "200px"}}]))
    assert LOGO.read_bytes() in _body_media(src)


def test_block_styles_print_in_the_docx(tmp_path):
    """Colour, size and background set on a block in the editor carry into the DOCX."""
    src = tmp_path / "t.docx"
    src.write_bytes(_compile([
        {"type": "heading", "attrs": {"level": 1, "textColor": "#FF8C00", "fontSize": "24pt", "textAlign": "center"},
         "content": [{"type": "text", "text": "Styled Title"}]},
        {"type": "paragraph", "attrs": {"backgroundColor": "#FEF9C3"}, "content": [{"type": "text", "text": "Shaded"}]},
    ]))
    doc = Document(str(src))
    title = next(p for p in doc.paragraphs if p.text == "Styled Title")
    run = title.runs[0]
    assert str(run.font.color.rgb) == "FF8C00"
    assert run.font.size.pt == 24
    shaded = next(p for p in doc.paragraphs if p.text == "Shaded")
    assert 'w:fill="FEF9C3"' in shaded._p.xml


def test_college_logo_block_prints_the_institute_crest(tmp_path):
    src = tmp_path / "t.docx"
    src.write_bytes(_compile([{"type": "imageBlock", "attrs": {"imageType": "college-logo", "src": ""}}]))
    assert Path("assets/svnit_crest.png").read_bytes() in _body_media(src)
