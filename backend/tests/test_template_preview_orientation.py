"""Regression test for Template Preview orientation and dimensions (Task A).

Verifies:
1. Every section in a compiled portrait template has WD_ORIENT.PORTRAIT, width ~ 8.27in, height ~ 11.69in.
2. Every section in a compiled landscape template has WD_ORIENT.LANDSCAPE, width ~ 11.69in, height ~ 8.27in.
3. The rendered page image through the preview pipeline is taller than it is wide for portrait, and wider than it is tall for landscape.
4. Dummy values are realistic instead of raw placeholders.
"""

from __future__ import annotations

import io
import tempfile
from pathlib import Path
import docx
from docx.enum.section import WD_ORIENT
import pymupdf
import yaml
from docxtpl import DocxTemplate

from server.app.core.template_compiler import (
    TemplateCompiler,
    _build_editor_dummy_context,
    safe_render_docx,
)
from docgen.renderer import Renderer
from docgen.schema import TemplateSchema


def test_template_preview_orientation_and_dimensions():
    for orientation in ("portrait", "landscape"):
        editor_json = {
            "type": "doc",
            "content": [
                {
                    "type": "heading",
                    "attrs": {"level": 1, "textAlign": "center"},
                    "content": [{"type": "text", "text": f"Orientation {orientation.title()}"}],
                },
                {
                    "type": "paragraph",
                    "content": [
                        {"type": "text", "text": "Official event for "},
                        {
                            "type": "fieldChip",
                            "attrs": {
                                "fieldName": "event_name",
                                "label": "Event Name",
                                "fieldType": "text",
                            },
                        },
                        {"type": "text", "text": " at "},
                        {
                            "type": "fieldChip",
                            "attrs": {
                                "fieldName": "venue",
                                "label": "Venue",
                                "fieldType": "text",
                            },
                        },
                        {"type": "text", "text": "."},
                    ],
                },
                {
                    "type": "signatureGrid",
                    "attrs": {
                        "columns": 2,
                        "signatures": [
                            {"label": "Student Head", "designation": "ACM SVNIT"},
                            {"label": "Faculty Advisor", "designation": "SVNIT"},
                        ],
                    },
                },
            ],
            "design_settings": {
                "orientation": orientation,
                "layout": "certificate" if orientation == "landscape" else "document",
            },
        }

        fields = [
            {"key": "event_name", "label": "Event Name", "type": "text", "required": True},
            {"key": "venue", "label": "Venue", "type": "text", "required": True},
        ]

        docx_bytes, schema_yaml, report = TemplateCompiler.compile(
            editor_json=editor_json,
            fields=fields,
            template_meta={
                "id": f"test_{orientation}",
                "name": f"Test {orientation}",
                "category": "Test",
                "description": "",
            },
            club_slug="acm",
            letterhead_cfg={"banner_placeholder": True},
            design_settings=editor_json["design_settings"],
        )

        # 1. Assert every section has the right orientation and page width/height
        doc = docx.Document(io.BytesIO(docx_bytes))
        assert len(doc.sections) > 0
        for section in doc.sections:
            if orientation == "landscape":
                assert section.orientation == WD_ORIENT.LANDSCAPE
                assert abs(section.page_width.inches - 11.69) < 0.05
                assert abs(section.page_height.inches - 8.27) < 0.05
            else:
                assert section.orientation == WD_ORIENT.PORTRAIT
                assert abs(section.page_width.inches - 8.27) < 0.05
                assert abs(section.page_height.inches - 11.69) < 0.05

        # 2. Render through the full preview pipeline to PDF & image
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
            tmp_dir = Path(td)
            tmpl_docx = tmp_dir / "template.docx"
            tmpl_docx.write_bytes(docx_bytes)
            schema_dict = yaml.safe_load(schema_yaml)
            schema_obj = TemplateSchema(**schema_dict)
            dummy_ctx = _build_editor_dummy_context(schema_obj, fields, editor_json)

            # Assert realistic dummy data used
            assert dummy_ctx["event_name"] == "Annual Photography Walk"
            assert dummy_ctx["venue"] == "Central Lawn, SVNIT"

            dtpl = DocxTemplate(str(tmpl_docx))
            safe_render_docx(dtpl, dummy_ctx, timeout=30.0, strict=False)
            rendered_docx = tmp_dir / "rendered.docx"
            dtpl.save(str(rendered_docx))

            renderer = Renderer()
            pdf_path, _ = renderer._convert_to_pdf(rendered_docx)
            pdf_doc = pymupdf.open(str(pdf_path))
            assert len(pdf_doc) > 0
            page = pdf_doc[0]
            page_scale = max(1400.0 / max(page.rect.width, 1.0), 2.0)
            pix = page.get_pixmap(matrix=pymupdf.Matrix(page_scale, page_scale))
            w, h = pix.width, pix.height
            pdf_doc.close()

            # 3. Assert rendered page image dimensions match orientation
            if orientation == "landscape":
                assert w > h, f"Landscape image must be wider than tall: got w={w}, h={h}"
            else:
                assert h > w, f"Portrait image must be taller than wide: got w={w}, h={h}"


def test_page_background_is_written_once():
    """_configure_page runs twice (init and after build); the DOCX must keep one <w:background>."""
    from docx.oxml.ns import qn

    editor_json = {
        "type": "doc",
        "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Tinted page"}]}],
    }
    design = {"orientation": "portrait", "pageBg": "#FFF8E7"}
    docx_bytes, _schema, _report = TemplateCompiler.compile(
        editor_json=editor_json,
        fields=[],
        template_meta={"id": "test_bg", "name": "Test bg", "category": "Test", "description": ""},
        club_slug="acm",
        letterhead_cfg={"banner_placeholder": True},
        design_settings=design,
    )
    doc = docx.Document(io.BytesIO(docx_bytes))
    assert len(doc.element.findall(qn("w:background"))) == 1


def test_dummy_values_match_whole_words():
    from server.app.core.template_compiler import _realistic_dummy_value

    assert _realistic_dummy_value("candidate_name") == "Sample Candidate Name"
    assert _realistic_dummy_value("student_count", "", "int") == 100
    assert _realistic_dummy_value("event_date", "", "date") == "15/10/2026"
    assert _realistic_dummy_value("venue") == "Central Lawn, SVNIT"
